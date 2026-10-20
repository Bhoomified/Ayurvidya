"""Bhoomi's Stage 4 tests: review sheet round-trip and backup.   python -m unittest tests.test_review"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from ayurvidya.review import export_review, import_review, review_summary  # noqa: E402
from ayurvidya.schema import Shloka, StepResult  # noqa: E402
from ayurvidya.store import Store  # noqa: E402


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.store = Store(self.d / "t.sqlite3")
        self.store.upsert_chapter("su.ni.1", "sushruta", "nidana", 1, "test")
        sh = Shloka(id="su.ni.1.3", samhita="sushruta", sthana="nidana", chapter=1, number="3", text="क ख ग")
        self.store.upsert_shlokas("su.ni.1", [sh])
        for step in (1, 2):
            self.store.save_step(StepResult("su.ni.1.3", step, f"out{step}", True, model="mock", prompt_version="v"))

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_round_trip(self):
        path = self.d / "review.xlsx"
        self.assertEqual(export_review(self.store, "su.ni.1", path), 2)
        wb = load_workbook(path)
        ws = wb.active
        head = [c.value for c in ws[1]]
        ws.cell(2, head.index("status") + 1, "approved")
        ws.cell(3, head.index("status") + 1, "reviewed")
        ws.cell(3, head.index("corrected_output") + 1, "fixed text")
        ws.cell(3, head.index("reviewer_note") + 1, "word order")
        wb.save(path)

        dry = import_review(self.store, path, dry_run=True)
        self.assertEqual(dry["updated"], 2)
        self.assertEqual(self.store.steps("su.ni.1.3")[1]["status"], "generated")   # dry run saved nothing

        r = import_review(self.store, path)
        self.assertEqual((r["updated"], r["corrected"], r["problems"]), (2, 1, []))
        steps = self.store.steps("su.ni.1.3")
        self.assertEqual(steps[1]["status"], "approved")
        self.assertEqual(steps[2]["output"], "fixed text")
        self.assertEqual(steps[2]["reviewer_note"], "word order")
        self.assertEqual(import_review(self.store, path)["unchanged"], 1)      # step 1 already approved
        self.assertEqual(review_summary(self.store, "su.ni.1")["approved"], 1)

    def test_bad_status_reported(self):
        path = self.d / "review.xlsx"
        export_review(self.store, "su.ni.1", path)
        wb = load_workbook(path)
        wb.active.cell(2, 9, "done")
        wb.save(path)
        self.assertEqual(len(import_review(self.store, path)["problems"]), 1)

    def test_backup(self):
        sys.path.insert(0, str(ROOT / "scripts"))
        import backup_db
        backup_db.backup_dir = lambda: self.d / "backups"
        dest = backup_db.backup(self.d / "t.sqlite3", keep=2)
        self.assertTrue(dest.exists())
        check = Store(dest)
        self.assertEqual(len(check.steps("su.ni.1.3")), 2)
        check.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)