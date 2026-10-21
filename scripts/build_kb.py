"""Build the knowledge base and load the shloka table.

    python scripts/build_kb.py                     # uses config/settings.yaml
    AYV_EMBEDDER=st:BAAI/bge-m3 python scripts/build_kb.py

Reads data/raw (mula), data/gold (validated tables) and every file in data/sources (*.txt, *.jsonl).
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")   # Devanagari on Windows consoles

from openpyxl import load_workbook  # noqa: E402

from ayurvidya.config import load_settings  # noqa: E402
from ayurvidya.index import KnowledgeBase  # noqa: E402
from ayurvidya.ingest import build_chunks, load_gold_xlsx, load_mula  # noqa: E402
from ayurvidya.store import Store  # noqa: E402


def topics_from_gold(paths) -> dict[str, str]:
    out = {}
    for p in paths:
        ws = load_workbook(p, read_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        h = [str(x or "") for x in rows[0]]
        if "Prakarana (topic)" in h:
            n, t = h.index("Verse No."), h.index("Prakarana (topic)")
            out.update({str(r[n]): str(r[t] or "") for r in rows[1:] if r[n]})
    return out

def topics_from_csv(path) -> dict[str, str]:
    """Topic names for verses not in the gold table (data/topics_su_ni_1.csv: verse,topic). Gold wins."""
    import csv
    if not path.exists():
        return {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {r["verse"].strip(): r["topic"].strip() for r in csv.DictReader(f) if r.get("topic")}
    

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default=None)
    ap.add_argument("--no-gold", action="store_true", help="build without few-shot gold examples (ablation)")
    args = ap.parse_args()
    s = load_settings(args.config)

    gold_files = sorted((s.data_dir / "gold").glob("*.xlsx"))
    topics = {**topics_from_csv(s.data_dir / "topics_su_ni_1.csv"), **topics_from_gold(gold_files)}
    shlokas, colophon = load_mula(s.data_dir / "raw" / "niimh_su_ni_1_mula.json", topics=topics)
    gold = {}
    if not args.no_gold:
        for g in gold_files:
            gold.update(load_gold_xlsx(g))
    sources = sorted(p for p in (s.data_dir / "sources").glob("*") if p.suffix in (".txt", ".jsonl"))
    chunks = build_chunks(shlokas, gold, sources)
    kb = KnowledgeBase.build(chunks, s.embedder, s.kb_dir, s.vector_store)

    store = Store(s.db_path)
    store.upsert_chapter("su.ni.1", "sushruta", "nidana", 1, "वातव्याधिनिदानम् (Vatavyadhi Nidana)")
    store.upsert_shlokas("su.ni.1", shlokas)
    store.close()

    print(f"shlokas: {len(shlokas)}  (colophon: {colophon})")
    print(f"chunks:  {len(kb.chunks)}  by type: {dict(Counter(c.source_type for c in kb.chunks))}")
    print(f"embedder: {s.embedder}  dim: {getattr(kb.dense, 'dim', '?')}  -> {s.kb_dir}")


if __name__ == "__main__":
    main()
