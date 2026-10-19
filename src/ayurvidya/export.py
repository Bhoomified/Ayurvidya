"""Export cached outputs to the project's Excel layout (one row per verse, one column per step),
so scholars can review them in the same format as the hand-made template."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .store import Store

HEAD = ["Verse No.", "Shloka", "Sampradanam", "Padavibhaga", "Anvaya", "Anvayartha",
        "Bhavartha", "Padakrityam", "Dhvanitartha", "Status", "Validator issues"]


def export_xlsx(store: Store, chapter_id: str, path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = chapter_id
    ws.append(HEAD)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F3E3C3")
    for sh in store.shlokas(chapter_id):
        steps = store.steps(sh["id"])
        statuses = {d["status"] for d in steps.values()} or {"—"}
        issues = "; ".join(f"S{n}: {', '.join(d['issues'])}" for n, d in steps.items() if d["issues"])
        ws.append([sh["number"], sh["text"]] + [steps.get(n, {}).get("output", "") for n in range(1, 8)]
                  + [",".join(sorted(statuses)), issues])
    widths = [9, 40, 34, 48, 36, 46, 50, 70, 50, 12, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return Path(path)
