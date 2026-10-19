"""Turn raw files into Shloka records and knowledge-base Chunks.

Inputs (all under data/):
  raw/niimh_su_ni_1_mula.json   {"<verse no>": [lines...]} — the NIIMH mula text, one key per verse number
  gold/*.xlsx                   scholar-validated 7-step tables (columns as in the project template)
  sources/*.txt                 reference texts with a small front-matter header (see load_source_file)
  sources/*.jsonl               one JSON object per line: {"text", "source_type", "work", "shloka_ids"?, ...}

Chunking rule (from the project design): never cut through a shloka. Mula = one chunk per verse unit;
commentary = one chunk per commented verse (split further only at danda boundaries if very long);
lexicon/dhatu entries = one chunk per headword.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, Iterator

from .schema import Chunk, Shloka, STEPS, SOURCE_TYPES
from .text import normalize, strip_verse_markers

COLOPHON_PREFIX = "इति "
# Verse units taught together (gadya sutras 1 and 2 form one pratijna unit in the template).
DEFAULT_JOINS = {("1", "2")}


def shloka_id(chapter: int, number: str, sthana: str = "ni", samhita: str = "su") -> str:
    return f"{samhita}.{sthana}.{chapter}.{number}"


def load_mula(path: str | Path, chapter: int = 1, joins=DEFAULT_JOINS,
              topics: dict[str, str] | None = None) -> tuple[list[Shloka], str]:
    """Load the NIIMH mula JSON. Returns (shlokas in order, colophon line)."""
    raw: dict[str, list[str]] = json.loads(Path(path).read_text(encoding="utf-8"))
    colophon = ""
    units: dict[str, list[str]] = {}
    for num in sorted(raw, key=int):
        lines = []
        for line in raw[num]:
            if line.strip().startswith(COLOPHON_PREFIX):
                colophon = normalize(line)
            else:
                lines.append(line)
        units[num] = lines
    joined = {b for _, b in joins}
    out: list[Shloka] = []
    for num in sorted(units, key=int):
        if num in joined:
            continue
        group = next((j for j in joins if j[0] == num), (num,))
        lines = [ln for n in group for ln in units[n]]
        main, variant = _split_variant(lines)
        number = "-".join(group)
        out.append(Shloka(
            id=shloka_id(chapter, number), samhita="sushruta", sthana="nidana", chapter=chapter,
            number=number, text=normalize("\n".join(main)), variant=normalize("\n".join(variant)),
            is_prose=len(group) > 1 or _looks_prose(main), topic=(topics or {}).get(number, ""),
        ))
    return out, colophon


def _split_variant(lines: list[str]) -> tuple[list[str], list[str]]:
    """NIIMH prints bracketed text two ways: (a) an alternate reading after the verse (verse 10) or an
    extra line (verse 68) -> kept apart as `variant`; (b) whole verses regarded as additions (40-41)
    -> when every line of the unit is inside brackets, the text itself is the verse (brackets dropped)."""
    flags, depth = [], 0
    for ln in lines:
        opened = "(" in ln
        depth += ln.count("(")
        flags.append(depth > 0 or opened)
        depth -= ln.count(")")
        depth = max(depth, 0)

    def clean(ln: str) -> str:
        ln = re.sub(r"\)\s*\|\s*$", "", ln.strip())      # ") |" printed after a closing variant
        return ln.replace("(", "").replace(")", "").strip()

    if all(flags) or not any(flags):
        return [clean(l) for l in lines], []
    main = [clean(l) for l, f in zip(lines, flags) if not f]
    variant = [clean(l) for l, f in zip(lines, flags) if f]
    return main, variant


def _looks_prose(lines: list[str]) -> bool:
    return len(lines) == 1 and "।" not in normalize(lines[0])[:-6]


# ------------------------------------------------------------------ gold annotations
GOLD_COLUMNS = {  # template column header -> step number (None = metadata)
    "Sampradanam": 1, "Padavibhaga": 2, "Anvaya": 3, "Anvayartha": 4,
    "Bhavartha": 5, "Padakrityam": 6, "Dhvanitartha": 7,
}


def load_gold_xlsx(path: str | Path, chapter: int = 1) -> dict[str, dict[int, str]]:
    """Read a 7-step table (the project's Excel layout). Returns {shloka_id: {step: text}}."""
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h or "").strip() for h in rows[0]]
    num_col = next(i for i, h in enumerate(header) if h.lower().startswith("verse"))
    step_cols = {GOLD_COLUMNS[h]: i for i, h in enumerate(header) if h in GOLD_COLUMNS}
    gold: dict[str, dict[int, str]] = {}
    for r in rows[1:]:
        if not r or r[num_col] is None:
            continue
        sid = shloka_id(chapter, str(r[num_col]).strip())
        gold[sid] = {step: str(r[col] or "").strip() for step, col in step_cols.items() if r[col]}
    return gold


# ------------------------------------------------------------------ reference sources
_HEADER = re.compile(r"^#\s*(\w+)\s*:\s*(.*)$")


def load_source_file(path: str | Path) -> Iterator[Chunk]:
    """Plain-text reference file. Header lines (top of file, or before any block) look like:

        # source_type: commentary
        # work: Nibandhasangraha (Dalhana)
        # shloka: 3            <- verse number(s) this block comments on, e.g. "3" or "21,22"
        <text ...>

    A new '# shloka:' line starts a new block, so one file can hold a whole chapter's commentary."""
    p = Path(path)
    meta: dict[str, str] = {}
    buf: list[str] = []
    block = 0

    def flush() -> Iterator[Chunk]:
        nonlocal block
        text = normalize("\n".join(buf))
        if not text:
            return
        st = meta.get("source_type", "commentary")
        if st not in SOURCE_TYPES:
            raise ValueError(f"{p.name}: unknown source_type '{st}' (allowed: {SOURCE_TYPES})")
        chapter = int(meta.get("chapter", 1))
        sids = [shloka_id(chapter, n.strip()) for n in meta.get("shloka", "").split(",") if n.strip()]
        for part_no, part in enumerate(_split_long(text)):
            block += 1
            yield Chunk(id=f"{p.stem}.{block}", text=part, source_type=st, work=meta.get("work", p.stem),
                        shloka_ids=sids, meta={"file": p.name, "part": part_no})

    for line in p.read_text(encoding="utf-8").splitlines():
        m = _HEADER.match(line)
        if m:
            key, val = m.group(1).lower(), m.group(2).strip()
            if key == "shloka" and buf:
                yield from flush()
                buf = []
            meta[key] = val
        else:
            buf.append(line)
    yield from flush()


def load_source_jsonl(path: str | Path) -> Iterator[Chunk]:
    p = Path(path)
    for n, line in enumerate(p.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        d = json.loads(line)
        st = d.pop("source_type")
        if st not in SOURCE_TYPES:
            raise ValueError(f"{p.name}:{n}: unknown source_type '{st}'")
        yield Chunk(id=d.pop("id", f"{p.stem}.{n}"), text=normalize(d.pop("text")), source_type=st,
                    work=d.pop("work", p.stem), shloka_ids=d.pop("shloka_ids", []), meta=d)


def _split_long(text: str, max_chars: int = 1200) -> list[str]:
    """Split only at danda/sentence boundaries, never mid-sentence."""
    if len(text) <= max_chars:
        return [text]
    sents = re.split(r"(?<=[।॥.])\s+", text)
    parts, cur = [], ""
    for s in sents:
        if cur and len(cur) + len(s) > max_chars:
            parts.append(cur.strip())
            cur = ""
        cur += s + " "
    if cur.strip():
        parts.append(cur.strip())
    return parts


# ------------------------------------------------------------------ build the KB
def build_chunks(shlokas: list[Shloka], gold: dict[str, dict[int, str]] | None = None,
                 source_paths: Iterable[str | Path] = ()) -> list[Chunk]:
    chunks: list[Chunk] = []
    for i, s in enumerate(shlokas):
        chunks.append(Chunk(
            id=f"mula.{s.id}", text=s.text + (f"\n(पाठान्तरम्) {s.variant}" if s.variant else ""),
            source_type="mula", work="Sushruta Samhita (NIIMH)", shloka_ids=[s.id],
            meta={"number": s.number, "order": i, "topic": s.topic, "search": strip_verse_markers(s.text)},
        ))
    for sid, steps in (gold or {}).items():
        for step, txt in steps.items():
            chunks.append(Chunk(id=f"gold.{sid}.{step}", text=txt, source_type="gold",
                                work="Validated 7-step table", shloka_ids=[sid],
                                meta={"step": step, "step_name": STEPS[step]}))
    for p in source_paths:
        p = Path(p)
        loader = load_source_jsonl if p.suffix == ".jsonl" else load_source_file
        chunks.extend(loader(p))
    ids = [c.id for c in chunks]
    dupes = {x for x in ids if ids.count(x) > 1}
    if dupes:
        raise ValueError(f"Duplicate chunk ids: {sorted(dupes)[:5]}")
    return chunks


def save_chunks(chunks: list[Chunk], path: str | Path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        for c in chunks:
            f.write(json.dumps(c.to_dict(), ensure_ascii=False) + "\n")


def load_chunks(path: str | Path) -> list[Chunk]:
    return [Chunk.from_dict(json.loads(l)) for l in Path(path).read_text(encoding="utf-8").splitlines() if l.strip()]
