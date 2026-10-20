"""Scholar review round-trip: database -> Excel review sheet -> scholar fills it -> back into the database.

One row per (verse, step). The scholar fills only the yellow columns:
    status            generated | reviewed | approved
    corrected_output  leave empty if the AI output is fine; otherwise write the corrected text
    reviewer_note     why it was changed / what is wrong
Approved rows are never overwritten by a later AI run (see Store.cached).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .schema import STEPS
from .store import Store

COLUMNS = ["shloka_id", "verse", "step", "step_name", "shloka_text", "ai_output", "valid", "issues",
           "status", "corrected_output", "reviewer_note"]
EDITABLE = {"status", "corrected_output", "reviewer_note"}
WIDTHS = [14, 7, 6, 14, 40, 70, 7, 30, 12, 70, 30]
STATUSES = ("generated", "reviewed", "approved")


def export_review(store: Store, chapter_id: str, path: str | Path, only_status: str | None = None,
                  verses: set[str] | None = None) -> int:
    """Write the review sheet. Returns the number of rows written."""
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(COLUMNS)
    head_fill, edit_fill = PatternFill("solid", fgColor="F3E3C3"), PatternFill("solid", fgColor="FFF4B0")
    for i, cell in enumerate(ws[1]):
        cell.font = Font(bold=True)
        cell.fill = edit_fill if COLUMNS[i] in EDITABLE else head_fill
    n = 0
    for sh in store.shlokas(chapter_id):
        if verses and sh["number"] not in verses and not (set(sh["number"].split("-")) & verses):
            continue
        for step, d in store.steps(sh["id"]).items():
            if only_status and d["status"] != only_status:
                continue
            ws.append([sh["id"], sh["number"], step, STEPS[step], sh["text"], d["output"],
                       "yes" if d["valid"] else "NO", "; ".join(d["issues"]), d["status"], "", d["reviewer_note"]])
            n += 1
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{max(n + 1, 2)}")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return n


def import_review(store: Store, path: str | Path, dry_run: bool = False) -> dict:
    """Read a filled review sheet and save status / corrections / notes. Returns counts + problems."""
    wb = load_workbook(path, read_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()                      # Windows: release the file, or it stays locked
    head = [str(h or "").strip() for h in rows[0]]
    missing = [c for c in ("shloka_id", "step", "status", "corrected_output", "reviewer_note") if c not in head]
    if missing:
        raise ValueError(f"{path}: missing columns {missing} - use a sheet made by export_review")
    col = {c: head.index(c) for c in head}
    out = {"updated": 0, "corrected": 0, "unchanged": 0, "problems": []}
    for n, r in enumerate(rows[1:], 2):
        if not r or not r[col["shloka_id"]]:
            continue
        sid, step = str(r[col["shloka_id"]]).strip(), int(r[col["step"]])
        status = str(r[col["status"]] or "").strip().lower()
        corrected = str(r[col["corrected_output"]] or "").strip()
        note = str(r[col["reviewer_note"]] or "").strip()
        current = store.steps(sid).get(step)
        if current is None:
            out["problems"].append(f"row {n}: {sid} step {step} is not in the database")
            continue
        if status not in STATUSES:
            out["problems"].append(f"row {n}: status '{status}' must be one of {STATUSES}")
            continue
        if status == current["status"] and not corrected and note == current["reviewer_note"]:
            out["unchanged"] += 1
            continue
        if not dry_run:
            store.set_review(sid, step, status, note, corrected or None)
        out["updated"] += 1
        out["corrected"] += bool(corrected)
    return out


def review_summary(store: Store, chapter_id: str) -> dict:
    """How many step outputs are generated / reviewed / approved, and how many verses are fully approved."""
    counts = {s: 0 for s in STATUSES}
    full = 0
    shl = store.shlokas(chapter_id)
    for sh in shl:
        steps = store.steps(sh["id"])
        for d in steps.values():
            counts[d["status"]] += 1
        full += len(steps) == 7 and all(d["status"] == "approved" for d in steps.values())
    return {"verses": len(shl), "verses_fully_approved": full, **counts}


__all__ = ["export_review", "import_review", "review_summary"]